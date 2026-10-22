# -*- coding: utf-8 -*-
import re, sys, importlib.util
from collections import Counter
from docx import Document
docx = sys.argv[1]
cont = sys.argv[2] if len(sys.argv) > 2 else None
d = Document(docx)
paras = [p.text for p in d.paragraphs]
cells = [c.text for t in d.tables for r in t.rows for c in r.cells]
all_t = '\n'.join(paras + cells)
spec = importlib.util.spec_from_file_location('m57q', cont)
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
pblocks = [b[1] for b in m.C if b[0] == 'p']
pchars = sum(len(t) for t in pblocks)
allchars = sum(len(t) for t in paras) + sum(len(c) for c in cells)
unres = sum(t.count('{R:') for t in paras + cells)
print('chars(all): %d chars(p): %d tables: %d unresolved: %d' % (allchars, pchars, len(d.tables), unres))
cites = re.findall(r'\[(\d+)(?:, [\d–\-]+-б\.)?\]', all_t)
used = sorted({int(x) for x in cites})
ir0 = next(i for i, t in enumerate(paras) if t.strip().startswith('РӮЙХАТИ'))
nref = len([t for t in paras[ir0:] if re.match(r'^\d+\. ', t.strip()) and len(t.strip()) > 40])
uncited = [n for n in range(1, nref + 1) if n not in used]
extra = [n for n in used if n < 1 or n > nref]
print('cites: %d refs: %d uncited: %s extra: %s' % (len(cites), nref, uncited, extra))
ir = next(i for i, t in enumerate(paras) if t.strip().startswith('РӮЙХАТИ'))
refl = [t for t in paras[ir:] if re.match(r'^\d+\. ', t.strip()) and len(t.strip()) > 40]
print('2020+: %d/%d URL: %d' % (sum(1 for t in refl if re.search(r'202[0-6]', t)), len(refl), sum(1 for t in refl if 'URL' in t)))
mix = sorted({w for t in paras + cells for w in re.findall(r'(?<![A-Za-z0-9.])(\d+)-(?:ро|ҳо|он)\b', t)})
print('mixed:', mix if mix else [])
caps = [b[1] for b in m.C if b[0] == 'cap' and b[1].startswith('Ҷадвали')]
capd = [k for k, v in Counter(caps).items() if v > 1]
print('caps: %d capdups: %s' % (len(caps), capd if capd else []))
nums = [int(re.match(r'Мисоли (\d+)', b[1]).group(1)) for b in m.C if b[0] == 'p' and b[1].startswith('Мисоли')]
print('examples: %d seq ok: %s' % (len(nums), nums == list(range(1, len(nums) + 1))))
sents = []
for t in pblocks:
    for x in re.split(r'(?<=[.!?])\s', t):
        x = x.strip()
        if len(x) > 25:
            sents.append(x)
sd = sum(v - 1 for v in Counter(sents).values() if v > 1)
print('HaminTarz: %d arab: %d vhead: %d sents-dups: %d Rleft: %d' % (len(re.findall(r'Ҳамин тарз|Ҳамин тавр|Ҳамин нуқта', all_t)), len(re.findall(r'[\u0600-\u06FF]', all_t)), sum(1 for c in cells if len(c.strip()) == 1 and not c.strip().isdigit()), sd, unres))
